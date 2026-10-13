"""Генерация вымышленных файлов для e2e и ручных проверок (KON-47): ``python qa/testdata/make_cv_files.py``.

Создаёт ``qa/testdata/cv/``: DOCX корпоративного формата (python-docx), версии для архива FR-039, PDF и PNG
для «Материалов», повреждённый и пустой DOCX, недопустимые форматы, файл больше лимита 10 МБ. Реальных
данных нет: имена, компании и стек вымышлены.
"""
import struct
import zlib
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent / "cv"
OUT.mkdir(exist_ok=True)


def cv(path: str, name: str, stack: str, years: int) -> None:
    d = Document()
    d.add_heading(name, level=1)
    for title, body in (
        ("Профиль", f"Системный аналитик, опыт работы {years} лет. Банковские интеграции, REST API."),
        ("Основной стек", stack),
        ("Отрасли", "Банки, страхование"),
        ("Проекты", "Платёжный шлюз: проектирование REST API, SQL-витрины. Стек: " + stack),
        ("Образование", "Высшее техническое"),
    ):
        d.add_heading(title, level=2)
        d.add_paragraph(body)
    d.save(OUT / path)


cv("cv_main_ok.docx", "Тестов Тест Тестович", "SQL, BPMN, UML, REST, Kafka, Jira", 7)
for v in range(2, 6):
    cv(f"cv_main_v{v}.docx", "Тестов Тест Тестович", f"SQL, BPMN, версия {v}", 7 + v)
cv("cv_extra.docx", "Тестов Тест Тестович", "Python, Camunda", 5)
(OUT / "materials.pdf").write_bytes(b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF\n")


def png() -> bytes:
    def chunk(t: bytes, d: bytes) -> bytes:
        c = struct.pack(">I", len(d)) + t + d
        return c + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"\x00\xff\x00\x00\xff"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw[:4])) + chunk(b"IEND", b""))


(OUT / "materials.png").write_bytes(png())
(OUT / "cv_corrupt.docx").write_bytes(b"PK\x03\x04" + b"garbage" * 50)
(OUT / "cv_empty.docx").write_bytes(b"")
(OUT / "cv_legacy_rtf.doc").write_bytes(b"{\rtf1\ansi Hello}")
(OUT / "cv_wrong.txt").write_text("not a cv", encoding="utf-8")
(OUT / "cv_wrong.exe").write_bytes(b"MZ\x90\x00")
(OUT / "cv_as_pdf.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
(OUT / "cv_big.docx").write_bytes(b"PK\x03\x04" + b"0" * (10 * 1024 * 1024 + 1024))
print("создано:", ", ".join(sorted(p.name for p in OUT.iterdir())))
