# -*- coding: utf-8 -*-
"""
Copyright (c) 2026 Konstantin Shakurov. Licensed under AGPL v3.
Commercial licensing: konstshakurov@gmail.com

Экспорт estimate.md/estimate_pp.md (отчёт скилла analyst_estimate) в xlsx
по шаблону импорта в CashMeter (1С). Правила маппинга колонок -
references/xlsx_export_rules.md рядом с этим скриптом.

Использование:
    python export_to_xlsx.py <путь_к_estimate.md> [--out <путь_к.xlsx>] [--stage "Этап 1"]

Если --out не задан - xlsx сохраняется рядом с md-файлом с тем же именем
(estimate.md -> estimate.xlsx). Единственный лист результата называется "Шаблон".
"""
import argparse
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

SCRIPT_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = SCRIPT_DIR.parent / "assets" / "Шаблон для импорта.xlsx"

PLAIN_FONT = Font(name="Arial", size=8, bold=False)
PLAIN_BORDER = Border()
PLAIN_FILL = PatternFill(fill_type=None)
PLAIN_ALIGNMENT = Alignment()

HEADER_ALIASES = {
    "№": ["№"],
    "Артефакт": ["Артефакт"],
    "Кол-во": ["Кол-во"],
    "Норма": ["Норма", "Ставка (ч./ед.)", "Ставка (ч.)"],
    "Итого": ["Итого", "Итого (ч.)"],
    "Описание": ["Описание"],
}


def parse_md_tables(md_path):
    """Разбирает markdown-файл на группы: [{heading, columns, rows: [{col: value}]}]
    Группа - это заголовок H3 (###), сразу за которым следует markdown-таблица."""
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    groups = []
    current_heading = None
    i, n = 0, len(lines)
    while i < n:
        line = lines[i].rstrip("\n")
        h3 = re.match(r"^###\s+(.*)$", line)
        if h3:
            current_heading = h3.group(1).strip()
            i += 1
            continue
        if line.strip().startswith("|") and current_heading is not None:
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            if len(table_lines) >= 2:
                header_cells = [c.strip() for c in table_lines[0].strip("|").split("|")]
                data_rows = []
                for row_line in table_lines[2:]:  # пропускаем заголовок и строку-разделитель
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    if len(cells) != len(header_cells):
                        continue
                    data_rows.append(dict(zip(header_cells, cells)))
                groups.append({"heading": current_heading, "columns": header_cells, "rows": data_rows})
                current_heading = None
            continue
        i += 1
    return groups


def resolve_col(columns, canonical):
    for alias in HEADER_ALIASES[canonical]:
        if alias in columns:
            return alias
    return None


def strip_group_prefix(heading):
    return re.sub(r"^(?:[A-ZА-Я]\.|\d+(?:\.\d+)*\.?)\s+", "", heading).strip()


def strip_md_bold(s):
    return re.sub(r"\*\*(.*?)\*\*", r"\1", s).strip()


def extract_count(raw):
    """Возвращает (значение, is_default). is_default=True значит, что реального
    числа в оценке не было (% / текст) и по правилу шаблона подставлена 1."""
    raw = strip_md_bold(raw)
    if "%" in raw:
        return 1, True
    m = re.match(r"^\s*(\d+)(?:[.,]\d+)?", raw)
    if not m:
        return 1, True
    return int(m.group(1)), False


def extract_hours(raw):
    raw = strip_md_bold(raw)
    m = re.search(r"(\d+(?:[.,]\d+)?)", raw)
    if not m:
        return None
    return float(m.group(1).replace(",", "."))


def extract_norma_hours(raw):
    """Пытается извлечь чистое число часов из столбца "Норма"/"Ставка".
    Возвращает (значение, fallback_needed). fallback_needed=True значит, что
    "Норма" - не чистое число (формула, диапазон, "фикс." и т.п.) и нужно
    брать столбец "Итого" вместо неё."""
    raw = strip_md_bold(raw).strip()
    cleaned = re.sub(r"\s*ч\.?\s*$", "", raw).strip()
    if re.match(r"^\d+(?:[.,]\d+)?$", cleaned):
        return float(cleaned.replace(",", ".")), False
    return None, True


def build_xlsx(md_path, out_path, stage_label="Этап 1"):
    groups = parse_md_tables(md_path)

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    for name in list(wb.sheetnames):
        if name != "ШАБЛОН":
            del wb[name]
    ws = wb["ШАБЛОН"]
    ws.title = "_tmp_rename_"  # openpyxl считает "ШАБЛОН"->"Шаблон" коллизией (регистронезависимо)
    ws.title = "Шаблон"

    row_idx = 2
    stats = {
        "groups_found": 0, "rows_written": 0, "rows_skipped": 0,
        "count_defaulted": 0, "descr_blank": 0, "hours_fallback_rows": [],
    }

    for g in groups:
        columns = g["columns"]
        col_no = resolve_col(columns, "№")
        col_art = resolve_col(columns, "Артефакт")
        col_qty = resolve_col(columns, "Кол-во")
        col_norma = resolve_col(columns, "Норма")
        col_total = resolve_col(columns, "Итого")
        col_descr = resolve_col(columns, "Описание")

        if not (col_no and col_art and col_total):
            continue  # это не таблица детализации (например, сводная таблица)

        stats["groups_found"] += 1
        task_name = strip_group_prefix(g["heading"])

        for row in g["rows"]:
            no_raw = strip_md_bold(row.get(col_no, ""))
            if not re.match(r"^\d+$", no_raw):
                stats["rows_skipped"] += 1
                continue  # строка "Подитого"/"Итого" и т.п.

            artefact = strip_md_bold(row.get(col_art, ""))
            if col_qty:
                qty, qty_is_default = extract_count(row.get(col_qty, ""))
            else:
                qty, qty_is_default = 1, True
            descr = strip_md_bold(row.get(col_descr, "")) if col_descr else ""

            hours = None
            fallback = True
            if col_norma:
                hours, fallback = extract_norma_hours(row.get(col_norma, ""))
            if fallback:
                hours = extract_hours(row.get(col_total, ""))
                stats["hours_fallback_rows"].append(
                    (task_name, artefact, row.get(col_norma, "") if col_norma else "(нет столбца)")
                )

            if qty_is_default:
                stats["count_defaulted"] += 1
            if not descr:
                stats["descr_blank"] += 1

            for col in range(1, 17):
                c = ws.cell(row=row_idx, column=col)
                c.font = copy(PLAIN_FONT)
                c.border = copy(PLAIN_BORDER)
                c.fill = copy(PLAIN_FILL)
                c.alignment = copy(PLAIN_ALIGNMENT)

            ws.cell(row=row_idx, column=1, value="Аналитика")            # A Фаза
            ws.cell(row=row_idx, column=2, value=stage_label)             # B Этап
            ws.cell(row=row_idx, column=3, value=task_name)               # C Задача
            ws.cell(row=row_idx, column=4, value=artefact)                # D Подзадача
            ws.cell(row=row_idx, column=5, value=qty)                     # E Кол-во задач
            ws.cell(row=row_idx, column=6, value="Аналитик")              # F Роль
            if hours is not None:
                c = ws.cell(row=row_idx, column=7, value=hours)           # G Часы min
                c.number_format = "0.00"
                c_max = ws.cell(row=row_idx, column=8, value=round(hours * 1.2, 2))  # H Часы max
                c_max.number_format = "0.00"
            if descr:
                ws.cell(row=row_idx, column=16, value=descr)              # P Комментарий

            row_idx += 1
            stats["rows_written"] += 1

    wb.save(out_path)
    return stats


def main():
    parser = argparse.ArgumentParser(description="Экспорт estimate.md в xlsx для импорта в CashMeter")
    parser.add_argument("estimate_md", help="Путь к estimate.md / estimate_pp.md")
    parser.add_argument("--out", help="Путь к выходному .xlsx (по умолчанию - рядом с md, то же имя)")
    parser.add_argument("--stage", default="Этап 1", help='Значение колонки "Наименование Этапа" (по умолчанию "Этап 1")')
    args = parser.parse_args()

    md_path = Path(args.estimate_md)
    out_path = Path(args.out) if args.out else md_path.with_suffix(".xlsx")

    stats = build_xlsx(md_path, out_path, stage_label=args.stage)
    fallback_rows = stats.pop("hours_fallback_rows")

    print(f"Сохранено: {out_path}")
    print(stats)
    if fallback_rows:
        print('Строки, где "Норма" не была чистым числом - взято "Итого":')
        for task, art, norma_raw in fallback_rows:
            print(f"   - [{task}] {art}: Норма={norma_raw!r}")


if __name__ == "__main__":
    main()
